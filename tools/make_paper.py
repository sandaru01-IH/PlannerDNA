"""Builds the PlannerDNA technical paper as PDF and DOCX from one content definition.

    .media-venv\\Scripts\\python.exe tools\\make_paper.py [--api http://127.0.0.1:8765]

Numbers come from the running loopback backend (active model, closed blinded tests, export).
Methodology diagrams and equations are drawn with matplotlib into deliverables/paper_assets/.
Output: deliverables/PlannerDNA_Technical_Paper.pdf and .docx
"""
from __future__ import annotations

import argparse
import html
import json
import re
import urllib.parse
import urllib.request
from collections import Counter
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyBboxPatch, Rectangle  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DELIV = ROOT / "deliverables"
FIGS = DELIV / "figures"
SHOTS = DELIV / "screenshots"
ASSETS = DELIV / "paper_assets"
WIN_FONTS = Path("C:/Windows/Fonts")

TITLE = ("PlannerDNA: Learning and Blind-Testing a Personal Pairwise Preference Model from Explicit "
         "Design Choices in a Synthetic 3D Urban-Planning Sandbox")
SHORT_TITLE = "PlannerDNA: explicit planning choices, an interpretable model and a blinded test"
AUTHOR = "Alpha"
AFFILIATION = "Independent Researcher"
DATE = "October 2026"

TEAL, TEAL_D, INK, GREY, ROSE = "#10b981", "#0f766e", "#0f172a", "#64748b", "#e11d48"

# ---------------------------------------------------------------- data


def get(api: str, path: str):
    if urllib.parse.urlparse(api).hostname not in {"127.0.0.1", "localhost", "::1"}:
        raise SystemExit("Loopback backends only.")
    with urllib.request.urlopen(api.rstrip("/") + path, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


def load(api: str) -> dict:
    cur = get(api, "/api/v1/models/current")["model"]
    ex = get(api, "/api/v1/export")
    plans = sorted((p for p in get(api, "/api/v1/evaluation")["plans"] if p["status"] == "closed"),
                   key=lambda p: p["closed_utc"])
    reports = {p["plan_id"]: get(api, f"/api/v1/evaluation/plans/{p['plan_id']}/report") for p in plans}
    run = next(r for r in ex["model_runs"] if r["id"] == cur["model_version"])
    ids = set(run["training_comparison_ids_json"])
    train = [c for c in ex["comparisons"] if c["comparison_id"] in ids]
    synthetic = lambda c: "synthetic" in (c.get("reason") or "").lower()  # noqa: E731

    def is_user_plan(pid: str) -> bool:
        tids = {t["comparison_id"] for t in reports[pid]["report"]["trials"] if t.get("comparison_id")}
        return not any(synthetic(c) for c in ex["comparisons"] if c["comparison_id"] in tids)

    user_plan = next(pid for pid in reports if is_user_plan(pid))
    agent_plan = next((pid for pid in reports if pid != user_plan), None)
    manifest = json.loads((FIGS / "figures_manifest.json").read_text(encoding="utf-8"))
    order = [e["plan_id"] for e in manifest["closed_evaluations"]]
    return {
        "model": cur, "export": ex, "train": train,
        "train_rule": sum(synthetic(c) for c in train),
        "per_scenario": Counter((c["scenario_id"], synthetic(c)) for c in train),
        "all_rule": sum(synthetic(c) for c in ex["comparisons"]),
        "modes": Counter((c["mode"], c["choice"], bool(c["training_eligible"])) for c in ex["comparisons"]),
        "user": reports[user_plan], "agent": reports.get(agent_plan) if agent_plan else None,
        "user_tag": f"_{order.index(user_plan) + 1}" if len(order) > 1 else "",
        "manifest": manifest,
    }


def pct(v: float, d: int = 1) -> str:
    return f"{v * 100:.{d}f}%"


# ---------------------------------------------------------------- diagrams


def _canvas(w_in: float, h_in: float):
    fig = plt.figure(figsize=(w_in, h_in), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100 * h_in / w_in)
    ax.set_aspect("equal")
    ax.axis("off")
    return fig, ax


def _box(ax, x, y, w, h, text, fc="#ecfdf5", ec=TEAL_D, fs=7.2, bold=False, ls="-", tc=INK, lw=1.0):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0,rounding_size=1.1", fc=fc, ec=ec, lw=lw, ls=ls))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs, color=tc,
            weight="bold" if bold else "normal", linespacing=1.25)


def _arrow(ax, x1, y1, x2, y2, color="#334155", ls="-", lw=1.0):
    ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                arrowprops=dict(arrowstyle="-|>,head_length=0.5,head_width=0.22", color=color, lw=lw, ls=ls,
                                shrinkA=0, shrinkB=0))


def _poly(ax, pts, color="#334155", ls="-"):
    xs, ys = zip(*pts[:-1])
    ax.plot(list(xs) + [pts[-2][0]], list(ys) + [pts[-2][1]], color=color, lw=1.0, ls=ls)
    _arrow(ax, *pts[-2], *pts[-1], color=color, ls=ls)


def _label(ax, x, y, text, fs=6.4, color=GREY, ha="center", style="italic"):
    ax.text(x, y, text, ha=ha, va="center", fontsize=fs, color=color, style=style, linespacing=1.2)


def diagram_architecture(path: Path) -> None:
    fig, ax = _canvas(7.2, 4.7)
    ax.add_patch(FancyBboxPatch((1.5, 2), 43, 61, boxstyle="round,pad=0,rounding_size=1.5", fc="#f8fafc", ec="#94a3b8", lw=1.0))
    ax.text(23, 60.3, "Unity 6 client (C#, URP, UI Toolkit)", ha="center", fontsize=8, weight="bold", color=INK)
    rows = [("UI views: Home, guided tour, compare,\nLearning Lab, blinded test", "#ffffff"),
            ("PlannerApp: selection, tools, overlays,\nsession and journey state", "#ffffff"),
            ("Pure C# core: immutable CityState, SplitMix64\ngenerator, metrics, validator, cost rules", "#ecfdf5"),
            ("CandidateGenerator (suggestions) and\nevaluation pair rule evalpair-1.0.0", "#ffffff"),
            ("LocalDecisionStore: JSON written first,\nfile-per-message ordered outbox", "#fef9c3"),
            ("SyncService, LearningClient, HealthMonitor\n(loopback HTTP only)", "#ffffff")]
    y, prev_bottom = 50.5, None
    for t, fc in rows:
        _box(ax, 4, y, 38, 7, t, fc=fc)
        if prev_bottom is not None:
            _arrow(ax, 23, prev_bottom, 23, y + 7)
        prev_bottom = y
        y -= 8.3
    net_mid = prev_bottom + 3.5
    ax.add_patch(FancyBboxPatch((54, 20), 44.5, 43, boxstyle="round,pad=0,rounding_size=1.5", fc="#f8fafc", ec="#94a3b8", lw=1.0))
    ax.text(76.25, 60.3, "Local backend (Python 3.12, FastAPI)", ha="center", fontsize=8, weight="bold", color=INK)
    brows = ["HTTP routes /api/v1, bound to 127.0.0.1",
             "records: validate, re-hash, recompute,\nidempotent insert",
             "models: readiness, train, predict-pair,\nrank-candidates",
             "evaluation: plans, sealed trials,\nreveal, close, report"]
    y, prev_bottom = 50.5, None
    for i, t in enumerate(brows):
        h = 6.6 if i else 5.6
        _box(ax, 56.5, y, 39.5, h, t, fc="#ffffff" if i else "#ecfdf5")
        if prev_bottom is not None:
            _arrow(ax, 76.25, prev_bottom, 76.25, y + h)
        prev_bottom = y
        y -= 8.3 if i else 7.6
    api_mid = 50.5 + 2.8
    _box(ax, 56.5, 10.5, 19, 7, "SQLite database\n(migrations, FK,\ntransactions)", fc="#e0f2fe", ec="#0369a1")
    _box(ax, 77, 10.5, 19, 7, "models/*.json\nweights + SHA-256\n(never pickle)", fc="#e0f2fe", ec="#0369a1")
    _arrow(ax, 66, 20, 66, 17.5)
    _arrow(ax, 86.5, 20, 86.5, 17.5)
    _box(ax, 56.5, 2, 39.5, 6.5, "GET /export \u2192 make_figures.py \u2192\nfigures, CSVs, manifest with refit check",
         fc="#f1f5f9", ec="#475569", fs=6.6)
    _poly(ax, [(42, net_mid), (49.2, net_mid), (49.2, api_mid), (56.5, api_mid)], color=TEAL_D)
    _arrow(ax, 45, net_mid, 42, net_mid, color=TEAL_D)
    ax.text(49.2, 33, "JSON over\n127.0.0.1\n\nidempotent:\nclient id +\npayload\nSHA-256", ha="center", va="center",
            fontsize=6.0, color=GREY, style="italic", linespacing=1.2,
            bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="none"))
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


def diagram_labelling(path: Path) -> None:
    fig, ax = _canvas(7.2, 3.3)
    top = [(2, "Explore\nseeded 20 \u00d7 20 city"), (21, "Edit\nvalidator + metrics"),
           (40, "Save A and B\nfull snapshots"), (59, "Compare\nsame baseline"), (78, "Choose\nA, B, tie or skip")]
    for i, (x, t) in enumerate(top):
        _box(ax, x, 33, 17, 8.5, t, fc="#ecfdf5" if i != 4 else "#fef9c3")
        if i:
            _arrow(ax, x - 2, 37.25, x, 37.25)
    _box(ax, 64, 14.5, 34, 12, "Eligibility gate (server re-checks)\nchoice is A or B  \u2022  mode = train\nboth designs feasible  \u2022  same scenario\nand baseline  \u2022  features-v1",
         fc="#ffffff", ec=TEAL_D, fs=6.8, lw=1.3)
    _arrow(ax, 86.5, 33, 86.5, 26.5)
    _label(ax, 93.2, 29.8, "outbox \u2192\nrecompute", fs=5.8)
    _box(ax, 33, 15.5, 24, 10, "Training set\nx = f(A) \u2212 f(B), y = 1[A]\n(server-recomputed features)", fc="#d1fae5", ec=TEAL_D, bold=False)
    _arrow(ax, 64, 20.5, 57, 20.5)
    _label(ax, 60.5, 23, "passes", fs=6.2)
    _box(ax, 64, 1.5, 34, 8, "Stored but never trained on:\nties, skips, infeasible pairs, evaluation mode", fc="#f1f5f9", ec="#94a3b8", fs=6.6)
    _arrow(ax, 81, 14.5, 81, 9.5)
    _label(ax, 85.5, 12, "fails", fs=6.2)
    _box(ax, 2, 15.5, 26, 10, "Edits, undo/redo, refused\nattempts, dwell time:\nlogged, never used as labels", fc="#f8fafc", ec="#94a3b8", ls="--", fs=6.6, tc="#334155")
    _poly(ax, [(29.5, 33), (29.5, 28.5), (15, 28.5), (15, 25.5)], color="#94a3b8", ls="--")
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


def diagram_training(path: Path) -> None:
    fig, ax = _canvas(7.2, 3.6)
    r1 = [(2, "Eligible comparisons\n(features recomputed\nby the server)", "#ecfdf5"),
          (26, "Differences\nx = f(A) \u2212 f(B)\nidentical outcomes dropped", "#ffffff"),
          (50, "Readiness\nn \u2265 10, scenarios \u2265 2,\nrank of X \u2265 2", "#fef9c3")]
    for i, (x, t, fc) in enumerate(r1):
        _box(ax, x, 37, 21.5, 10, t, fc=fc)
        if i:
            _arrow(ax, x - 2.5, 42, x, 42)
    _box(ax, 76, 37, 22, 10, "Refuse (HTTP 422)\nreasons shown,\nnothing written", fc="#ffe4e6", ec=ROSE)
    _arrow(ax, 71.5, 42, 76, 42)
    _label(ax, 73.8, 44.5, "fails", fs=6.2)
    r2 = [(2, "Mirror each row\n(x, y) and (\u2212x, 1 \u2212 y)", "#ffffff"),
          (26, "Scale by RMS\n(no centring keeps\nthe symmetry)", "#ffffff"),
          (50, "L2 logistic regression\nno intercept, C = 1,\nlbfgs, fixed seed", "#d1fae5"),
          (76, "Model file (JSON)\nversion = time + data\nfingerprint; SHA-256", "#e0f2fe")]
    for i, (x, t, fc) in enumerate(r2):
        _box(ax, x, 19.5, 21.5 if i < 3 else 22, 10, t, fc=fc)
        if i:
            _arrow(ax, r2[i - 1][0] + 21.5, 24.5, x, 24.5)
    _poly(ax, [(60.75, 37), (60.75, 33.2), (12.75, 33.2), (12.75, 29.5)])
    _label(ax, 36.5, 35, "passes", fs=6.2)
    _box(ax, 18, 2, 30, 10, "Leave-one-scenario-out CV\naccuracy, Brier, log loss vs\nchance, cheaper, more access,\nequal weights", fc="#f8fafc", ec="#475569", fs=6.5)
    _box(ax, 54, 2, 30, 10, "Scenario bootstrap (B = 200)\n90% percentile intervals,\nsign agreement per weight", fc="#f8fafc", ec="#475569", fs=6.5)
    _arrow(ax, 57, 19.5, 40, 12)
    _arrow(ax, 64, 19.5, 69, 12)
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


def diagram_protocol(path: Path) -> None:
    fig, ax = _canvas(7.2, 7.8)
    H = 100 * 7.8 / 7.2
    lanes = {"P": 12, "U": 44, "B": 82}
    for key, name in (("P", "Participant"), ("U", "Unity client"), ("B", "Backend (server)")):
        _box(ax, lanes[key] - 10, H - 7, 20, 5, name, fc="#0f172a", ec="#0f172a", tc="white", bold=True, fs=7.5)
        ax.plot([lanes[key], lanes[key]], [2, H - 7], color="#94a3b8", lw=0.9, ls=(0, (3, 3)))

    def msg(y, a, b, text, dashed=False):
        x1, x2 = lanes[a], lanes[b]
        _arrow(ax, x1, y, x2, y, color="#0f172a", ls="--" if dashed else "-")
        ax.text((x1 + x2) / 2, y + 1.3, text, ha="center", va="bottom", fontsize=6.3, color=INK)

    def note(y, lane, w, text, fc="#ecfdf5", ec=TEAL_D, h=5.6):
        x = lanes[lane] - w / 2
        x = min(max(x, 1), 99 - w)
        _box(ax, x, y - h / 2, w, h, text, fc=fc, ec=ec, fs=6.1)

    y = H - 12.5
    msg(y, "U", "B", "1  register plan: expected model, unseen scenario seeds, pairs, repeats")
    y -= 8
    note(y, "B", 36, "2  freeze model (version + file SHA-256)\ndraw plan_seed, candidate_seed_base (secrets)\nstore registration SHA-256; training locked", h=7.2)
    y -= 7.5
    loop_top = y
    y -= 4.5
    msg(y, "B", "U", "3  next: scenario seed, candidate seed", dashed=True)
    y -= 6.5
    note(y, "U", 40, "4  generate candidates, pick a trade-off pair\n(evalpair-1.0.0; the rule never sees the model)", fc="#fef9c3", ec="#a16207")
    y -= 6.5
    msg(y, "U", "B", "5  register trial (designs A, B)")
    y -= 7
    note(y, "B", 33, "6  recompute features; p_A = \u03c3(w \u00b7 x)\nseal h = SHA-256(payload with nonce)", h=6.2)
    y -= 7
    msg(y, "B", "U", "7  seal h and screen sides (never p_A)", dashed=True)
    y -= 6
    msg(y, "U", "P", "8  Left / Right designs, no prediction", dashed=True)
    y -= 5.5
    msg(y, "P", "U", "9  choice")
    y -= 5.5
    msg(y, "U", "B", "10  evaluation-mode preference (never training data)")
    y -= 3.2
    ax.add_patch(Rectangle((2, y), 96, loop_top - y, fill=False, ec="#64748b", lw=0.9))
    ax.add_patch(Rectangle((2, loop_top - 3.4), 46, 3.4, fc="#e2e8f0", ec="#64748b", lw=0.9))
    ax.text(3, loop_top - 1.7, "loop: every trial (first presentations, then repeats, sides swapped)", fontsize=6.0, va="center", color=INK)
    y -= 5
    msg(y, "U", "B", "11  Finish: reveal every answered trial")
    y -= 5.5
    msg(y, "B", "U", "sealed payloads; Unity checks SHA-256(payload) = h", dashed=True)
    y -= 5.5
    msg(y, "U", "B", "12  close: report stored as canonical JSON + SHA-256")
    y -= 5.5
    msg(y, "B", "U", "report, re-hashed and recomputed on every read", dashed=True)
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


def diagram_walk(path: Path) -> None:
    n = 7
    fig = plt.figure(figsize=(3.4, 3.55), dpi=300)
    ax = fig.add_axes([0.02, 0.1, 0.96, 0.88])
    ax.set_xlim(0, n)
    ax.set_ylim(n, 0)
    ax.set_aspect("equal")
    ax.axis("off")
    road = lambda r, c: r in (0, 3, 6) or c == 3  # noqa: E731
    park, h_ok, h_no = (1, 5), (4, 1), (5, 1)
    for r in range(n):
        for c in range(n):
            fc = "#cbd5e1" if road(r, c) else "#f8fafc"
            if (r, c) == park:
                fc = "#4ade80"
            elif (r, c) == h_ok:
                fc = "#14b8a6"
            elif (r, c) == h_no:
                fc = "#fb7185"
            ax.add_patch(Rectangle((c, r), 1, 1, fc=fc, ec="white", lw=1.2))
    ax.text(park[1] + 0.5, park[0] + 0.5, "park", ha="center", va="center", fontsize=6.5, weight="bold")
    ax.text(h_ok[1] + 0.5, h_ok[0] + 0.5, "H1", ha="center", va="center", fontsize=7, weight="bold", color="white")
    ax.text(h_no[1] + 0.5, h_no[0] + 0.5, "H2", ha="center", va="center", fontsize=7, weight="bold", color="white")
    p1 = [(3, 1), (3, 2), (3, 3), (2, 3), (1, 3), (0, 3), (0, 4), (0, 5)]
    p2 = [(6, 1), (6, 2), (6, 3), (5, 3), (4, 3), (3, 3), (2, 3), (1, 3), (0, 3), (0, 4), (0, 5)]
    for pts, col, off in ((p1, TEAL_D, -0.12), (p2, ROSE, 0.12)):
        ax.plot([c + 0.5 + off for r, c in pts], [r + 0.5 + off for r, c in pts], color=col, lw=1.8, marker="o", ms=2.4)
    fig.text(0.5, 0.055, "H1: 40 + 7 \u00d7 40 + 40 = 360 m  (accessible)", ha="center", fontsize=7, color=TEAL_D)
    fig.text(0.5, 0.015, "H2: 40 + 10 \u00d7 40 + 40 = 480 m  (not accessible)", ha="center", fontsize=7, color=ROSE)
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------- equations


def render_eq(latex: str, path: Path, size: float = 11.5) -> tuple[float, float]:
    plt.rcParams["mathtext.fontset"] = "stix"
    fig = plt.figure(figsize=(0.01, 0.01), dpi=300)
    fig.text(0, 0, f"${latex}$", fontsize=size)
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.03, transparent=False, facecolor="white")
    plt.close(fig)
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    return w / 300.0, h / 300.0  # inches


# ---------------------------------------------------------------- references

REFS = {
    "who2016": "WHO Regional Office for Europe, <i>Urban Green Spaces and Health: A Review of Evidence</i>. Copenhagen, Denmark: World Health Organization, 2016.",
    "geertman2004": "S. Geertman and J. Stillwell, \u201cPlanning support systems: An inventory of current practice,\u201d <i>Computers, Environment and Urban Systems</i>, vol. 28, no. 4, pp. 291\u2013310, 2004.",
    "poplin2012": "A. Poplin, \u201cPlayful public participation in urban planning: A case study for online serious games,\u201d <i>Computers, Environment and Urban Systems</i>, vol. 36, no. 3, pp. 195\u2013206, 2012.",
    "thurstone1927": "L. L. Thurstone, \u201cA law of comparative judgment,\u201d <i>Psychological Review</i>, vol. 34, no. 4, pp. 273\u2013286, 1927.",
    "bradley1952": "R. A. Bradley and M. E. Terry, \u201cRank analysis of incomplete block designs: I. The method of paired comparisons,\u201d <i>Biometrika</i>, vol. 39, no. 3/4, pp. 324\u2013345, 1952.",
    "luce1959": "R. D. Luce, <i>Individual Choice Behavior: A Theoretical Analysis</i>. New York, NY, USA: Wiley, 1959.",
    "mcfadden1974": "D. McFadden, \u201cConditional logit analysis of qualitative choice behavior,\u201d in <i>Frontiers in Econometrics</i>, P. Zarembka, Ed. New York, NY, USA: Academic Press, 1974, pp. 105\u2013142.",
    "louviere2000": "J. J. Louviere, D. A. Hensher, and J. D. Swait, <i>Stated Choice Methods: Analysis and Applications</i>. Cambridge, U.K.: Cambridge University Press, 2000.",
    "herbrich2000": "R. Herbrich, T. Graepel, and K. Obermayer, \u201cLarge margin rank boundaries for ordinal regression,\u201d in <i>Advances in Large Margin Classifiers</i>, A. J. Smola, P. Bartlett, B. Sch\u00f6lkopf, and D. Schuurmans, Eds. Cambridge, MA, USA: MIT Press, 2000, pp. 115\u2013132.",
    "joachims2002": "T. Joachims, \u201cOptimizing search engines using clickthrough data,\u201d in <i>Proc. 8th ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining</i>, 2002, pp. 133\u2013142.",
    "furnkranz2010": "J. F\u00fcrnkranz and E. H\u00fcllermeier, Eds., <i>Preference Learning</i>. Berlin, Germany: Springer, 2010.",
    "christiano2017": "P. F. Christiano, J. Leike, T. B. Brown, M. Martic, S. Legg, and D. Amodei, \u201cDeep reinforcement learning from human preferences,\u201d in <i>Advances in Neural Information Processing Systems 30</i>, 2017, pp. 4299\u20134307.",
    "kaufman2012": "S. Kaufman, S. Rosset, C. Perlich, and O. Stitelman, \u201cLeakage in data mining: Formulation, detection, and avoidance,\u201d <i>ACM Transactions on Knowledge Discovery from Data</i>, vol. 6, no. 4, art. 15, 2012.",
    "kapoor2023": "S. Kapoor and A. Narayanan, \u201cLeakage and the reproducibility crisis in machine-learning-based science,\u201d <i>Patterns</i>, vol. 4, no. 9, art. 100804, 2023.",
    "tversky1974": "A. Tversky and D. Kahneman, \u201cJudgment under uncertainty: Heuristics and biases,\u201d <i>Science</i>, vol. 185, no. 4157, pp. 1124\u20131131, 1974.",
    "nosek2018": "B. A. Nosek, C. R. Ebersole, A. C. DeHaven, and D. T. Mellor, \u201cThe preregistration revolution,\u201d <i>Proceedings of the National Academy of Sciences</i>, vol. 115, no. 11, pp. 2600\u20132606, 2018.",
    "oliver2007": "L. N. Oliver, N. Schuurman, and A. W. Hall, \u201cComparing circular and network buffers to examine the influence of land use on walking for leisure and errands,\u201d <i>International Journal of Health Geographics</i>, vol. 6, art. 41, 2007.",
    "samuelson1938": "P. A. Samuelson, \u201cA note on the pure theory of consumer\u2019s behaviour,\u201d <i>Economica</i>, vol. 5, no. 17, pp. 61\u201371, 1938.",
    "tversky1969": "A. Tversky, \u201cIntransitivity of preferences,\u201d <i>Psychological Review</i>, vol. 76, no. 1, pp. 31\u201348, 1969.",
    "rudin2019": "C. Rudin, \u201cStop explaining black box machine learning models for high stakes decisions and use interpretable models instead,\u201d <i>Nature Machine Intelligence</i>, vol. 1, no. 5, pp. 206\u2013215, 2019.",
    "hoerl1970": "A. E. Hoerl and R. W. Kennard, \u201cRidge regression: Biased estimation for nonorthogonal problems,\u201d <i>Technometrics</i>, vol. 12, no. 1, pp. 55\u201367, 1970.",
    "roberts2017": "D. R. Roberts <i>et al.</i>, \u201cCross-validation strategies for data with temporal, spatial, hierarchical, or phylogenetic structure,\u201d <i>Ecography</i>, vol. 40, no. 8, pp. 913\u2013929, 2017.",
    "efron1979": "B. Efron, \u201cBootstrap methods: Another look at the jackknife,\u201d <i>The Annals of Statistics</i>, vol. 7, no. 1, pp. 1\u201326, 1979.",
    "efron1993": "B. Efron and R. J. Tibshirani, <i>An Introduction to the Bootstrap</i>. New York, NY, USA: Chapman &amp; Hall, 1993.",
    "field2007": "C. A. Field and A. H. Welsh, \u201cBootstrapping clustered data,\u201d <i>Journal of the Royal Statistical Society: Series B</i>, vol. 69, no. 3, pp. 369\u2013390, 2007.",
    "brier1950": "G. W. Brier, \u201cVerification of forecasts expressed in terms of probability,\u201d <i>Monthly Weather Review</i>, vol. 78, no. 1, pp. 1\u20133, 1950.",
    "gneiting2007": "T. Gneiting and A. E. Raftery, \u201cStrictly proper scoring rules, prediction, and estimation,\u201d <i>Journal of the American Statistical Association</i>, vol. 102, no. 477, pp. 359\u2013378, 2007.",
    "niculescu2005": "A. Niculescu-Mizil and R. Caruana, \u201cPredicting good probabilities with supervised learning,\u201d in <i>Proc. 22nd Int. Conf. Machine Learning (ICML)</i>, 2005, pp. 625\u2013632.",
    "naeini2015": "M. P. Naeini, G. F. Cooper, and M. Hauskrecht, \u201cObtaining well calibrated probabilities using Bayesian binning,\u201d in <i>Proc. 29th AAAI Conf. Artificial Intelligence</i>, 2015, pp. 2901\u20132907.",
    "wilson1927": "E. B. Wilson, \u201cProbable inference, the law of succession, and statistical inference,\u201d <i>Journal of the American Statistical Association</i>, vol. 22, no. 158, pp. 209\u2013212, 1927.",
    "agresti1998": "A. Agresti and B. A. Coull, \u201cApproximate is better than \u2018exact\u2019 for interval estimation of binomial proportions,\u201d <i>The American Statistician</i>, vol. 52, no. 2, pp. 119\u2013126, 1998.",
    "mcnemar1947": "Q. McNemar, \u201cNote on the sampling error of the difference between correlated proportions or percentages,\u201d <i>Psychometrika</i>, vol. 12, no. 2, pp. 153\u2013157, 1947.",
    "sandve2013": "G. K. Sandve, A. Nekrutenko, J. Taylor, and E. Hovig, \u201cTen simple rules for reproducible computational research,\u201d <i>PLoS Computational Biology</i>, vol. 9, no. 10, art. e1003285, 2013.",
    "fips180": "National Institute of Standards and Technology, \u201cSecure Hash Standard (SHS),\u201d FIPS PUB 180-4, Gaithersburg, MD, USA, Aug. 2015.",
    "steele2014": "G. L. Steele Jr., D. Lea, and C. H. Flood, \u201cFast splittable pseudorandom number generators,\u201d in <i>Proc. ACM Int. Conf. Object Oriented Programming Systems Languages &amp; Applications (OOPSLA)</i>, 2014, pp. 453\u2013472.",
    "kleppmann2019": "M. Kleppmann, A. Wiggins, P. van Hardenberg, and M. McGranaghan, \u201cLocal-first software: You own your data, in spite of the cloud,\u201d in <i>Proc. ACM SIGPLAN Int. Symp. New Ideas, New Paradigms, and Reflections on Programming and Software (Onward!)</i>, 2019, pp. 154\u2013178.",
    "pedregosa2011": "F. Pedregosa <i>et al.</i>, \u201cScikit-learn: Machine learning in Python,\u201d <i>Journal of Machine Learning Research</i>, vol. 12, pp. 2825\u20132830, 2011.",
    "harris2020": "C. R. Harris <i>et al.</i>, \u201cArray programming with NumPy,\u201d <i>Nature</i>, vol. 585, pp. 357\u2013362, 2020.",
    "virtanen2020": "P. Virtanen <i>et al.</i>, \u201cSciPy 1.0: Fundamental algorithms for scientific computing in Python,\u201d <i>Nature Methods</i>, vol. 17, pp. 261\u2013272, 2020.",
    "hunter2007": "J. D. Hunter, \u201cMatplotlib: A 2D graphics environment,\u201d <i>Computing in Science &amp; Engineering</i>, vol. 9, no. 3, pp. 90\u201395, 2007.",
    "unity": "Unity Technologies, \u201cUnity 6 real-time development platform,\u201d version 6000.6.7f1, software, 2026. [Online]. Available: https://unity.com",
    "fastapi": "S. Ram\u00edrez, \u201cFastAPI,\u201d software. [Online]. Available: https://fastapi.tiangolo.com",
    "sqlite": "D. R. Hipp <i>et al.</i>, \u201cSQLite,\u201d software. [Online]. Available: https://www.sqlite.org",
}

# ---------------------------------------------------------------- content


def build_content(d: dict, eqdir: Path) -> tuple[str, list[str], list[tuple]]:
    m = d["model"]
    coef = {c["feature"]: c for c in m["coefficients"]}
    diag = m["diagnostics"]
    lrule = {b["name"]: b for b in diag["baselines"]}
    u = d["user"]["report"]
    uver = d["user"]["verification"]
    um = u["metrics"]["model"]
    ub = {b["name"]: b for b in u["metrics"]["baselines"]}
    uc, upos, urc, ucal = u["counts"], u["position"], u["repeat_consistency"], u["calibration"]
    plan = u["plan"]
    ag = d["agent"]["report"] if d["agent"] else None
    tag = d["user_tag"]
    ps = d["per_scenario"]
    n_train = len(d["train"])
    n_rule = d["train_rule"]
    n_user = n_train - n_rule
    n_all = len(d["export"]["comparisons"])
    scored = [t for t in u["trials"] if t["scored"]]
    agree_rule = sum(1 for t in scored if (t["p_a"] > 0.5) == (t["baseline_more_access_p_a"] > 0.5)
                     and t["baseline_more_access_p_a"] != 0.5)
    pmin = min(max(t["p_a"], 1 - t["p_a"]) for t in scored)
    pv = m["package_versions"]
    acc = coef["delta_access_fraction"]

    abstract = (
        "Planning students are routinely asked to balance competing objectives, yet the trade-offs they actually "
        "favour are rarely made explicit, let alone tested. We present PlannerDNA, a local, single-user 3D sandbox in "
        "which a participant edits a procedurally generated 20 \u00d7 20 neighbourhood, saves two complete alternatives "
        "and records an explicit preference between them. Five transparent outcome metrics (park access within a "
        "400 m road-network walk, housing capacity, park share, intervention cost and lost original housing) are "
        "recomputed deterministically after every edit and independently re-verified by a local server. On request, "
        "a pairwise logistic (Bradley\u2013Terry) model with mirrored rows, no intercept and L2 regularisation is "
        "fitted to the participant's own feasible A/B choices, with scenario-level bootstrap intervals and "
        "leave-one-scenario-out diagnostics against simple decision rules. Predictive validity is then assessed by a "
        "pre-registered, blinded protocol: the model is frozen by checksum, test scenarios are verified unseen, every "
        "prediction is committed with a SHA-256 seal before the choice, screen sides are randomised from a "
        "server-held seed, repeated pairs measure the participant's own consistency, and the report is stored as "
        "hashed canonical JSON that is recomputed on every read. In a single-participant case study "
        f"({n_train} training comparisons across {m['scenario_count']} synthetic cities, {n_rule} of them generated "
        f"by a fixed rule), the frozen model predicted {um['correct']:g} of {um['n']} scored held-out choices "
        f"({pct(um['accuracy'])}; 95% scenario-bootstrap interval {pct(um['accuracy_ci'][0], 0)}\u2013"
        f"{pct(um['accuracy_ci'][1], 0)}), exactly matching the rule \u201calways choose more park access\u201d, with "
        f"which it agreed on every scored pair. We argue that the contribution is the workflow (explicit trade-offs, "
        f"explicit labels and an evaluation that cannot be flattered) rather than the accuracy figure.")
    keywords = ["preference learning", "paired comparisons", "Bradley\u2013Terry model", "planning support systems",
                "serious games", "blinded evaluation", "pre-registration", "reproducibility",
                "green-space accessibility"]

    E = {}

    def eq(key: str, latex: str):
        p = eqdir / f"eq_{key}.png"
        E[key] = (p, render_eq(latex, p))
        return ("eq", key, p, E[key][1])

    B: list[tuple] = []
    add = B.append

    # 1 Introduction
    add(("h1", "Introduction", "intro"))
    add(("p", "Urban planning is, to a large degree, the negotiation of trade-offs. Adding a park can bring more "
              "homes within walking distance of green space, which a substantial body of evidence associates with "
              "health and well-being [@who2016]; but the same land might otherwise hold housing, and every "
              "intervention has a cost and may displace existing residents. Planning support systems and serious "
              "games have long been proposed to make such trade-offs tangible for professionals, students and "
              "the public [@geertman2004; @poplin2012]. What they rarely do is record, explicitly and verifiably, "
              "which trade-offs a particular person actually prefers, and then test whether a summary of those "
              "preferences predicts that person's later choices."))
    add(("p", "Learning from pairwise comparisons is a mature field, from Thurstone's law of comparative judgment "
              "and the Bradley\u2013Terry and Luce models to conditional logit choice models [@thurstone1927; "
              "@bradley1952; @luce1959; @mcfadden1974], and from pairwise ranking in machine learning to "
              "preference-based reinforcement learning [@herbrich2000; @joachims2002; @furnkranz2010; "
              "@christiano2017]. With few choices, however, fitted preference weights are easily over-interpreted, "
              "and evaluations are easily flattered: by leakage between training and test data [@kaufman2012; "
              "@kapoor2023], by selecting test items that suit the model, by showing predictions that anchor the "
              "participant [@tversky1974], or by choosing the analysis after seeing the data [@nosek2018]."))
    add(("p", "This paper describes PlannerDNA, a system that couples a deterministic urban-design sandbox with an "
              "explicit labelling regime, a deliberately simple and interpretable preference model, and a blinded, "
              "pre-registered evaluation protocol whose integrity is enforced by the software rather than by "
              "convention. Its contributions are:"))
    add(("bullets", [
        "a deterministic, verifiable synthetic planning environment with five transparent proxy metrics, "
        "implemented in C# and re-implemented in Python, with every stored design recomputed by the server "
        "(#sec:env#);",
        "an explicit-choice labelling regime in which only feasible A/B choices made in training mode can become "
        "training data, and interaction traces never do (#sec:data#);",
        "a symmetric pairwise logistic model with readiness refusals, scenario-level bootstrap uncertainty and "
        "leave-one-scenario-out diagnostics against rules of thumb (#sec:model#);",
        "a blinded held-out evaluation protocol using cryptographic commitments, server-held randomisation, "
        "model-blind pair selection and repeat-consistency measurement, with a self-verifying report "
        "(#sec:eval#);",
        "a single-participant case study, reported together with the reasons it does not show more than a "
        "one-line rule (#sec:results#, #sec:discussion#)."]))

    # 2 Related work
    add(("h1", "Related work", "related"))
    add(("h2", "Planning support, serious games and accessibility measures", None))
    add(("p", "Planning support systems combine models, data and visualisation to inform planning tasks "
              "[@geertman2004], and serious games have been used to make participation more engaging and to "
              "expose trade-offs to non-experts [@poplin2012]. Accessibility to green space is commonly "
              "operationalised as the share of residents within a distance threshold of a park, with network "
              "(street) distance preferred over straight-line buffers because it reflects the routes people can "
              "actually walk [@oliver2007]. PlannerDNA adopts a road-network distance with a 400 m threshold, but "
              "its metrics are deliberately simplified proxies on a synthetic grid and are not calibrated to any "
              "real place."))
    add(("h2", "Paired comparisons and preference learning", None))
    add(("p", "The Bradley\u2013Terry model expresses the probability that one item is preferred to another as a "
              "logistic function of a difference in latent utilities [@bradley1952]; Thurstone's earlier "
              "formulation uses a normal link [@thurstone1927], and Luce's choice axiom generalises it to choice "
              "sets [@luce1959]. When utilities are linear in observed attributes, the model becomes a conditional "
              "logit, the workhorse of discrete-choice and stated-choice analysis [@mcfadden1974; @louviere2000]. "
              "In machine learning, the same \u201cpairwise transform\u201d (learning a linear classifier on attribute "
              "differences) underlies ranking SVMs [@herbrich2000; @joachims2002] and the broader preference "
              "learning literature [@furnkranz2010], and pairwise human judgments now train large reward models "
              "[@christiano2017]. PlannerDNA's labels are <i>stated</i> choices in a hypothetical setting, not "
              "revealed preferences in the economic sense [@samuelson1938], and human choices can be inconsistent "
              "or intransitive [@tversky1969], which is why the protocol measures repeat consistency."))
    add(("h2", "Interpretable models and evaluation integrity", None))
    add(("p", "For decisions that people must understand and contest, inherently interpretable models are preferable "
              "to post hoc explanations of opaque ones [@rudin2019]; a ridge-penalised linear model "
              "[@hoerl1970] on five named outcomes is about as interpretable as a predictive model can be. On the "
              "evaluation side, leakage is a leading cause of irreproducible machine-learning results "
              "[@kaufman2012; @kapoor2023]; grouped (blocked) cross-validation is recommended when observations "
              "are clustered [@roberts2017], and the cluster bootstrap resamples whole groups for the same reason "
              "[@efron1979; @efron1993; @field2007]. Pre-registration separates confirmatory from exploratory "
              "analysis [@nosek2018]. Probabilistic predictions are scored with proper scoring rules such as the "
              "Brier score and log loss [@brier1950; @gneiting2007], and calibration is assessed separately "
              "[@niculescu2005; @naeini2015]. Cryptographic hash functions [@fips180] allow a prediction to be "
              "committed before an outcome is known and verified afterwards. Finally, the local-first software "
              "movement argues for applications that keep working, and keep the user's data, without a cloud "
              "service [@kleppmann2019]."))

    # 3 System overview
    add(("h1", "System overview", "overview"))
    add(("p", "PlannerDNA consists of a Unity 6 client [@unity] and an optional local backend written in Python with "
              "FastAPI and SQLite [@fastapi; @sqlite] (#fig:arch#). Both run on one computer and communicate only "
              "over the loopback interface; there are no cloud services and no telemetry. The client's core (city "
              "state, generation, validation, metrics, cost, constraints, history and storage formats) is plain C# "
              "with no engine dependencies, so it is covered by fast unit tests. The city state is immutable: every "
              "accepted edit produces a new state, which makes undo and redo exact snapshot restores."))
    add(("fig", "arch", None, "Architecture. The Unity client writes every record locally first and delivers it, in "
         "order, through a file-per-message outbox; the backend re-hashes and recomputes every design before "
         "storing it. Model files are JSON with a SHA-256 checksum. Figures are generated only from an export.", 1.0))
    add(("p", "The client is offline-first in the sense of [@kleppmann2019]: the sandbox never waits for the server. "
              "Records are delivered strictly in order (so a decision can never arrive before the designs it refers "
              "to) and idempotently: each carries a client id and the SHA-256 of its payload, so a resend after a lost "
              "acknowledgement is a harmless duplicate, while reuse of an id with different content is refused."))
    add(("p", "A session follows four stages shown on a journey line at the top of the screen (#fig:ui#): "
              "<b>Explore</b> the generated neighbourhood; <b>Design A and B</b> from the same baseline; "
              "<b>Compare and choose</b>; and <b>Learn</b>, where the participant may train a model, inspect it, "
              "request ranked suggestions, or start a blinded test. A Home screen and a 13-step guided tour, written "
              "for first-year urban-planning undergraduates, spotlight each part of the interface and wait for the "
              "participant to perform the real action."))
    add(("fig", "ui", SHOTS / "07_design_a.png", "The sandbox. Left: tools and history. Centre: the 3D city with the "
         "park-access overlay (teal homes reach a park within 400 m by road; rose homes do not). Right: the mission, "
         "hard rules as green or red dots, the selected parcel and live metrics. Bottom: saved alternatives.", 1.0))

    # 4 Environment
    add(("h1", "Synthetic planning environment", "env"))
    add(("p", "The exercise is the <i>Neighbourhood Green-Space Dilemma</i>: increase the share of housing capacity "
              "within a 400 m road-network walk of a park while keeping housing capacity at or above 90% of the "
              "baseline, spending at most 100 synthetic credits and never exceeding eight floors. The mission text "
              "is generated from the scenario parameters (#tab:params#), so text and rules cannot drift apart."))
    add(("table", "params", "Scenario parameters (generator green-gen-1.0.0, metrics metrics-1.0.0, cost rules cost-rules-1.0.0).",
         ["Parameter", "Value"],
         [["Grid", "20 \u00d7 20 parcels of 40 m (800 m \u00d7 800 m)"],
          ["Budget", "100 synthetic credits"],
          ["Minimum housing", "ceil(0.9 \u00d7 baseline units), integer arithmetic"],
          ["Maximum height", "8 floors"],
          ["Access threshold", "400 m along roads"],
          ["Units per floor; floors of a new building", "4; 2"]], [0.42, 0.58]))
    add(("h2", "Deterministic generation", None))
    add(("p", "A seed fully determines the city. Generation uses SplitMix64 [@steele2014] rather than a platform "
              "random-number generator because its output is specified and identical across runtimes; a unit test "
              "pins reference values. Every third row is a road and three vertical roads are drawn from three bands, "
              "so every parcel touches a road. About 10% of parcels are vacant, commercial parcels concentrate beside "
              "the middle vertical road, residential parcels have two to six floors (35% zoned residential-only), "
              "one or two small parks sit in one quadrant, and one or two civic parcels are protected. A candidate "
              "city is accepted only if the road network is connected, baseline access is below 75%, there are at "
              "least three vacant parcels and at least one feasible vacant-to-park edit improves access; otherwise "
              "generation retries, up to 32 attempts, and then reports an error instead of producing a broken city."))
    add(("h2", "Metrics", None))
    add(("p", "Housing capacity is the sum over residential parcels of floors \u00d7 4 units, a capacity proxy rather "
              "than a population. Park access uses a walk distance along roads only (#fig:walk#): a multi-source "
              "breadth-first search starts from every road cell adjacent to a park, and a residential parcel uses its "
              "best adjacent road cell. With k road steps between the two road cells, the walk distance and access "
              "share are"))
    add(eq("walk", r"d_i = 40 + 40\,k_i + 40\ \mathrm{m}, \qquad i\ \mathrm{is\ served} \;\Leftrightarrow\; d_i \leq 400\ \mathrm{m}"))
    add(eq("access", r"\mathrm{Access} = 100 \cdot \dfrac{\sum_{i \in R} u_i\,\mathbf{1}[d_i \leq 400]}{\sum_{i \in R} u_i}, \qquad u_i = 4\,\mathrm{floors}_i"))
    add(("p", "where R is the set of residential parcels; access is reported as not applicable when housing is zero. "
              "With 40 m cells the threshold falls between 400 m (served) and the next possible distance, 440 m "
              "(not served). Park share is the number of park parcels over non-road parcels. Intervention cost compares "
              "each parcel's baseline state b<sub>i</sub> with its current state s<sub>i</sub>:"))
    add(eq("cost", r"\mathrm{Cost} = \sum_{i} c(b_i, s_i), \qquad c(b, b) = 0"))
    add(("p", "using the table in #tab:cost#. Because cost depends only on the difference between baseline and "
              "current state, restoring a parcel always costs nothing and the order of edits cannot game the budget. "
              "Lost original housing sums, over parcels that were residential at baseline, the units removed; it is "
              "a displacement proxy reported separately from net housing change."))
    add(("fig", "walk", None, "Road-network walk distance. Only road cells (grey) are walkable, with no diagonal moves "
         "or shortcuts through parcels. H1 reaches the park in 360 m and is served; H2, one parcel away, needs 480 m "
         "and is not.", 0.5))
    add(("table", "cost", "Cost rules (cost-rules-1.0.0); costs compare baseline and current parcel states.",
         ["Baseline \u2192 current", "Credits"],
         [["unchanged", "0"],
          ["vacant \u2192 park", "20"],
          ["vacant or park \u2192 new residential or commercial", "25 + 5 \u00d7 floors (35 for 2 floors)"],
          ["residential or commercial \u2192 park", "35"],
          ["residential or commercial \u2192 vacant (demolish)", "15"],
          ["park \u2192 vacant", "10"],
          ["residential \u2194 commercial", "20, + 5 per added floor or + 3 per removed floor"],
          ["same building use, floors changed", "+ 5 per added floor, + 3 per removed floor"],
          ["anything involving a road or civic parcel", "undefined: the edit is refused"]], [0.55, 0.45]))
    add(("h2", "Constraints, validation and cross-implementation checks", None))
    add(("p", "Four hard constraints are evaluated after every edit: budget (spent \u2264 100), housing floor, maximum "
              "floors, and a defined cost rule for every changed parcel. A validator refuses edits to roads and "
              "civic parcels, parks on parks, commercial uses on residential-only land and heights outside one to "
              "eight floors, each with a plain-language reason; refused attempts are counted but never enter the "
              "design timeline. Infeasible designs may be explored and saved, but are marked as such."))
    add(("p", "The metric engine exists twice: in C# for the interactive client and in Python for the server. Both "
              "compute the same FNV-1a content hash of a city. Contract fixtures exported from Unity are replayed by "
              "the Python test suite, and at run time the server rebuilds every submitted design from its snapshot, "
              "recomputes hash, metrics, constraints and features, and refuses any disagreement. The database "
              "therefore cannot hold a design whose numbers disagree with the published rules."))

    # 5 Data
    add(("h1", "Data collection and labelling", "data"))
    add(("p", "The unit of analysis is one explicit comparison between two saved designs. Edits, undo clicks, refused "
              "attempts and dwell time are logged for auditability but are never used as labels (#fig:label#). A "
              "comparison is <b>training-eligible</b> only if the choice is A or B (ties and skips are stored but not "
              "used), it was made in training mode (answers given during a blinded test are never training data, not "
              "even after the test closes), both designs are feasible, both come from the same registered scenario "
              "and baseline, and both carry the same feature version. Unity and the server both check these rules; "
              "the server's result is authoritative."))
    add(("fig", "label", None, "From interaction to label. Only an explicit A-or-B choice between two feasible designs "
         "made in training mode becomes a training example; everything else is stored for transparency but never "
         "trained on.", 1.0))
    add(("p", "Each design D is summarised by five features (features-v1), each a difference from the scenario "
              "baseline and oriented so that higher means \u201cmore of a commonly desirable outcome\u201d "
              "(#tab:features#). The orientation fixes only the sign convention; it is not a preference weight, and "
              "no weights exist until a model is trained on the participant's own choices."))
    add(("table", "features", "Feature vector f(D) (features-v1). H<sub>0</sub> is baseline housing capacity; all values are fractions.",
         ["Feature", "Definition", "Higher means"],
         [["delta_access_fraction", "(access % \u2212 baseline access %) / 100", "more homes within a 400 m walk of a park"],
          ["delta_housing_fraction", "(units \u2212 H<sub>0</sub>) / H<sub>0</sub>", "more housing capacity"],
          ["delta_park_share_fraction", "(park % \u2212 baseline park %) / 100", "more park land"],
          ["negative_budget_fraction", "\u2212 spent credits / 100", "lower cost"],
          ["negative_lost_housing_fraction", "\u2212 lost original units / H<sub>0</sub>", "fewer original homes removed"]],
         [0.33, 0.37, 0.30]))
    add(("p", "The model input for a comparison is the difference x = f(A) \u2212 f(B) with label y = 1 if A was "
              "chosen. Comparisons whose two designs have identical measured outcomes (x = 0) carry no information "
              "about trade-offs; they are counted, reported as ignored and never fitted."))

    # 6 Model
    add(("h1", "Preference model", "model"))
    add(("h2", "Specification and symmetry", None))
    add(("p", "The model is a linear Bradley\u2013Terry, or conditional logit, model on outcome differences "
              "[@bradley1952; @mcfadden1974]:"))
    add(eq("model", r"P(A \succ B \mid x) = \sigma(w^{\top}x) = \dfrac{1}{1 + e^{-w^{\top}x}}, \qquad x = f(A) - f(B)"))
    add(("p", "There is no intercept. An intercept would absorb any tendency to choose whichever design is presented "
              "as A, i.e. a screen-position habit, and would make the model's predictions depend on presentation "
              "order. Without it, and because \u03c3(\u2212z) = 1 \u2212 \u03c3(z),"))
    add(eq("sym", r"P(B \succ A) = \sigma(-w^{\top}x) = 1 - \sigma(w^{\top}x) = 1 - P(A \succ B)"))
    add(("p", "holds exactly. To make the fitted weights themselves independent of which design was labelled A, every "
              "comparison is added twice, as (x, y) and as (\u2212x, 1 \u2212 y). The mirrored copies balance the fit; "
              "they are not additional evidence, and every count reported anywhere is of unique comparisons."))
    add(("h2", "Estimation", None))
    add(("p", "Features are scaled by their root-mean-square over the 2n mirrored training rows. Because the "
              "mirrored rows have mean exactly zero, the standard deviation computed by a non-centring "
              "StandardScaler equals the RMS, and not centring preserves the symmetry. With every difference divided "
              "by its feature scale s<sub>j</sub>, the standardised weights solve an L2-penalised logistic regression "
              "[@hoerl1970]:"))
    add(eq("obj", r"\hat{w}^{(s)} = \mathrm{arg\,min}_{w}\ \ \dfrac{1}{2}\Vert w\Vert^2 + C\sum_{r=1}^{2n}\ell\left(y_r,\ \sigma(w^{\top}\tilde{x}_r)\right), \qquad s_j = \sqrt{\dfrac{1}{2n}\sum_{r=1}^{2n}x_{rj}^2}"))
    add(("p", f"where \u2113 is the logistic loss. The implementation uses scikit-learn's LogisticRegression "
              f"[@pedregosa2011] with C = 1, the lbfgs solver, a tolerance of 10<sup>\u22128</sup>, at most 1000 "
              f"iterations and a fixed random seed (20261025), on NumPy and SciPy [@harris2020; @virtanen2020]. A "
              f"standardised weight is the change in log-odds for one typical (RMS) training difference in that "
              f"feature, so standardised weights are comparable across features; raw weights w<sub>j</sub> = "
              f"w<sup>(s)</sup><sub>j</sub> / s<sub>j</sub> are used for utilities and probabilities. A feature "
              f"that never varied in training receives weight 0 and is reported as <i>not estimated</i> rather than "
              f"unimportant, and |standardised weight| below 0.1 (odds change of less than about 10%) is reported as "
              f"little or no association. The utility of a design, w \u00b7 f(D), is unit-free and meaningful only "
              f"for ordering designs within a scenario."))
    add(("h2", "Readiness", None))
    add(("p", "Training refuses rather than guessing. It requires at least 10 informative comparisons from at least "
              "two scenarios, and the RMS-scaled difference matrix must have rank at least 2, so that the data span at "
              "least two independent trade-off directions; otherwise the server returns HTTP 422 with the reason "
              "(not enough comparisons, a single scenario, or homogeneous comparisons) and writes nothing. Below 30 "
              "comparisons over five scenarios the model is labelled <i>exploratory</i> wherever it appears. These "
              "thresholds are heuristics, not statistical adequacy guarantees."))
    add(("h2", "Uncertainty and diagnostics", None))
    add(("p", "Choices within one scenario share a city and are not independent, so uncertainty is estimated by a "
              "scenario-level (cluster) bootstrap [@efron1993; @field2007]: scenarios are resampled with replacement "
              "200 times with the feature scale held fixed, the model is refitted, and each standardised weight gets "
              "a 90% percentile interval and the share of resamples whose sign agrees with the full fit. A weight "
              "whose interval includes zero is reported as directionally uncertain."))
    add(("p", "Leave-one-scenario-out (LOSO) cross-validation [@roberts2017] predicts each scenario's comparisons with "
              "a model and scaler fitted without that scenario. The same comparisons are scored by four transparent "
              "baselines: chance (always 0.5); always the cheaper design; always the design with more park access; "
              "and the sign of the equal-weight sum of the five features. Deterministic rules predict 1, 0 or 0.5 "
              "(for ties). With predictions p<sub>i</sub> and outcomes y<sub>i</sub>, accuracy counts an undecided "
              "prediction as half correct, and the Brier score [@brier1950] and log loss are the usual proper "
              "scoring rules [@gneiting2007]:"))
    add(eq("acc", r"\mathrm{Acc} = \dfrac{1}{n}\sum_{i=1}^{n}\left(\mathbf{1}\left[(p_i - \frac{1}{2})(y_i - \frac{1}{2}) > 0\right] + \frac{1}{2}\,\mathbf{1}\left[p_i = \frac{1}{2}\right]\right)"))
    add(eq("scores", r"\mathrm{BS} = \dfrac{1}{n}\sum_{i}(p_i - y_i)^2, \qquad \mathrm{LL} = -\dfrac{1}{n}\sum_{i}\left[y_i\log p_i + (1 - y_i)\log(1 - p_i)\right]"))
    add(("p", "Log loss is reported only for probabilistic predictors. Training-set accuracy is also shown, labelled "
              "optimistic, with a separability warning at 100%. LOSO is a diagnostic inside the training data and "
              "is never described as a held-out test; only a closed blinded evaluation (#sec:eval#) is."))
    add(("fig", "train", None, "Model fitting. Readiness checks gate training; mirrored, RMS-scaled differences are "
         "fitted by an L2 logistic regression without intercept; the result is a checksummed JSON file. Diagnostics "
         "are computed alongside and shown with the weights.", 1.0))
    add(("h2", "Training, versioning and use", None))
    add(("p", "Training happens only when the participant presses <b>Train my model</b>, and only once every local "
              "record has been delivered, so a model always reflects every choice made so far. The model version is "
              "<i>pm-&lt;UTC timestamp&gt;-&lt;fingerprint&gt;</i>, where the fingerprint is the SHA-256 of the "
              "canonical JSON of the ML and feature versions, hyperparameters and every training row; retraining on "
              "unchanged data returns the existing model. Each model is stored as a JSON file (never pickle, so "
              "loading cannot execute code) containing feature names, raw and standardised weights, the scale, "
              "hyperparameters, seed, training comparison and scenario ids, counts, fingerprint and package "
              "versions. Its SHA-256 is stored in the database, and every prediction reloads the file and refuses if "
              "the hash differs."))
    add(("p", "A trained model is used in three ways. In <b>Compare</b>, its prediction for the pair is requested only "
              "after the participant has confirmed a choice, to avoid anchoring [@tversky1974], and is shown with a "
              "note on whether that comparison was in the training data (#fig:compare#). In the <b>Learning Lab</b>, "
              "weights, intervals and diagnostics are displayed with plain-language statements. For "
              "<b>suggestions</b>, Unity generates up to 30 candidate designs by applying one to four seeded, "
              "validator-approved edits to the baseline (at most 800 attempts), drops infeasible designs, exact "
              "duplicates and near-duplicates, and sends them to the server, which recomputes each candidate, "
              "refuses tampered or infeasible ones, ranks them by utility and returns the top k. Suggestions are "
              "labelled as ranked synthetic candidates, never as optimal plans."))
    add(("fig", "compare", SHOTS / "10_compare_prediction.png", "Compare and choose. Both designs are tabulated "
         "against the same baseline with a plain-language trade-off summary. The participant chooses A, B, tie or "
         "skip and may give a reason; only then does the model's prediction appear.", 1.0))

    # 7 Evaluation protocol
    add(("h1", "Blinded held-out evaluation protocol", "eval"))
    add(("p", "The evaluation asks: <i>does the frozen model predict choices it has never seen, in cities it was never "
              "trained on, better than simple rules of thumb?</i> The protocol (blinded-eval-1.0.0, #fig:protocol#) is "
              "designed so that neither the software nor the participant can flatter the answer, and each safeguard is "
              "enforced by the server as a refusal rather than left to good practice."))
    add(("fig", "protocol", None, "Sequence of the blinded evaluation. The prediction for each trial is computed and "
         "committed before the trial is shown; the participant sees only Left and Right; all seals are opened and "
         "checked only after the last choice.", 0.98))
    add(("h2", "Pre-registration and freezing", None))
    add(("p", "The client registers a plan containing the expected model version, two to twenty evaluation scenarios "
              "(identifier and seed), one to ten pairs per scenario, a number of repeated pairs and the pair rule. The "
              "server freezes the model by storing its version and the SHA-256 of its file; every later prediction "
              "reloads the file and refuses if the hash differs. Each scenario is refused if its identifier or seed "
              "appears in the model's training data, if it already has any training-mode choice, or if an earlier plan "
              "used it; Unity draws fresh seeds between 10<sup>5</sup> and 10<sup>9</sup> and retries on refusal. The "
              "server then draws two secrets with a cryptographic random source, a plan seed (screen sides and repeat "
              "selection) and a candidate-seed base (which designs are generated), and stores the SHA-256 of the "
              "canonical registration. While a plan is open the server refuses to train any model, to accept practice "
              "choices on evaluation scenarios, to predict or rank on them, or to open a second plan (#tab:locks#). "
              "Exports taken mid-test omit unrevealed predictions."))
    add(("table", "locks", "Server-enforced refusals while a blinded test is open.",
         ["Attempt", "Response"],
         [["Train any model", "409 evaluation_in_progress"],
          ["Training-mode choice on an evaluation scenario", "409 scenario_reserved_for_evaluation"],
          ["predict-pair or rank-candidates on an evaluation scenario", "409 sealed_for_evaluation"],
          ["Open a second plan", "409 evaluation_already_open"],
          ["Scenario seen in training, practice or an earlier plan", "422 scenario_not_unseen"],
          ["Reveal a prediction before the choice is recorded", "409 choice_not_recorded"]], [0.58, 0.42]))
    add(("h2", "Model-blind pair selection", None))
    add(("p", "For trial t the candidate seed is the secret base plus t. Unity generates candidates with the same "
              "generator used for suggestions and, using a deterministic random generator seeded by the candidate "
              "seed, picks a feasible pair that is a genuine trade-off: each design is better on at least one of the "
              "five features (if no such pair exists, any pair with different features). Because this rule "
              "(evalpair-1.0.0) never sees the model, test pairs cannot be selected to suit it. The server refuses "
              "wrong seeds, wrong rules, infeasible, identical or uninformative pairs, and designs used in any earlier "
              "comparison or trial."))
    add(("h2", "Sealed predictions", None))
    add(("p", "On registration of a trial the server recomputes both designs' features, computes p<sub>A</sub> with "
              "the frozen weights and stores a sealed payload together with its hash:"))
    add(eq("seal", r"h_t = \mathrm{SHA256}\left(\mathrm{canon}\{t,\ \mathrm{plan},\ \mathrm{model},\ A,\ B,\ p_A,\ \nu_t\}\right), \qquad \nu_t:\ 128\ \mathrm{random\ bits}"))
    add(("p", "Only h<sub>t</sub> and the screen order are returned; p<sub>A</sub> is never sent before the choice. The "
              "random nonce \u03bd<sub>t</sub> prevents recovering p<sub>A</sub> by hashing candidate values, so the "
              "commitment is hiding; the collision resistance of SHA-256 [@fips180] makes it binding, so the server "
              "cannot later substitute a different prediction. The reveal endpoint works only after the choice has "
              "reached the server and returns the payload, so anyone can check that its hash equals the seal shown "
              "before the choice."))
    add(("h2", "Presentation, repeats and reveal", None))
    add(("p", "For first presentations, the left side shows A when a pseudo-random draw seeded by the plan seed and "
              "trial index falls below 0.5; the participant sees <i>Left</i> and <i>Right</i>, never A and B. After all "
              "first presentations, a server-seeded sample of pairs is shown again with sides swapped, as separate "
              "trials, to measure the participant's own consistency. Answers are ordinary preference records with "
              "mode = evaluation, the trial id, the locked model version and the side shown on the left; they are "
              "never training-eligible. Unity reveals nothing during the test: showing each prediction straight "
              "after each choice would let the participant learn what the model expects and drift towards it, "
              "inflating accuracy. Only when the participant presses <b>Finish</b> does Unity request every reveal, "
              "re-hash each payload, check it against its seal, trial and pair, and close the plan (#fig:evalui#)."))
    add(("fig", "evalui", SHOTS / "16_eval_reveals.png", "End of a blinded test in the client: every sealed prediction "
         "has been opened and checked against its fingerprint before the report is shown. (Frame from the scripted "
         "demonstration run, not the case-study test.)", 1.0))
    add(("h2", "Scoring and the self-verifying report", None))
    add(("p", "Only first presentations answered A or B are scored; ties, skips, unanswered trials and repeats are "
              "counted but not scored, so the denominator is always stated. The report gives accuracy, Brier score "
              "and log loss for the model, and the same measures for chance, cheaper, more park access, equal weights "
              "and <i>always the left design</i>, which detects a side habit, each with the model-minus-baseline "
              "accuracy difference. Intervals come from 1000 bootstrap resamples of whole scenarios (seed 20261025) "
              "with 95% percentile limits, omitted when fewer than three scenarios are scored. Repeat consistency is "
              "reported with a Wilson score interval [@wilson1927; @agresti1998]:"))
    add(eq("wilson", r"\dfrac{\hat{p} + \frac{z^2}{2n} \pm z\sqrt{\frac{\hat{p}(1 - \hat{p})}{n} + \frac{z^2}{4n^2}}}{1 + \frac{z^2}{n}}"))
    add(("p", "The share of left choices is tested against 50% with an exact two-sided binomial test, and "
              "calibration bins the probability of the favoured design (0.5 to 1.0 in steps of 0.1) and reports the "
              "expected calibration error [@naeini2015]:"))
    add(eq("ece", r"\mathrm{ECE} = \sum_{b} \dfrac{n_b}{n}\left|\bar{c}_b - \bar{o}_b\right|"))
    add(("p", "where, for bin b with n<sub>b</sub> choices, the two terms are the mean confidence and the observed "
              "rate at which the favoured design was chosen. A "
              "held-out guarantee recomputes the number of evaluation choices and evaluation scenarios present in the "
              "frozen model's training data (zero by construction). The report is computed once at closing, stored as "
              "canonical JSON with its SHA-256, and on every read the server both re-hashes the stored report and "
              "recomputes it from the stored rows; it is marked reproducible only when all three hashes agree. "
              "Warnings are generated for small samples, fewer than five scenarios, intervals that include 50%, a "
              "baseline at least as accurate as the model, low consistency and side bias."))

    # 8 Implementation
    add(("h1", "Implementation and verification", "impl"))
    add(("p", "#tab:stack# summarises the implementation. Every version identifier (generator, metrics, cost rules, "
              "features, schema, ML and protocol) is written into each stored record, so any number can be traced to the "
              "rules that produced it."))
    add(("table", "stack", "Implementation stack.",
         ["Component", "Technology"],
         [["Client", "Unity 6000.6.7f1, Universal Render Pipeline, UI Toolkit, Input System; C#"],
          ["Server", f"Python {pv['python']}, FastAPI, Pydantic, SQLite (migrations, foreign keys, BEGIN IMMEDIATE)"],
          ["Model", f"scikit-learn {pv['scikit-learn']}, NumPy {pv['numpy']}, SciPy {pv['scipy']}"],
          ["Figures", "matplotlib [@hunter2007] from export files; one CSV per figure; manifest with refit check"],
          ["Determinism", "SplitMix64 generation; FNV-1a content hashes; SHA-256 payloads, models and reports"],
          ["Network", "loopback only (127.0.0.1); no cloud services, no telemetry"]], [0.22, 0.78]))
    add(("p", "Verification is automated (#tab:tests#). Tests that need a server start a throwaway backend on a "
              "separate port with a temporary data directory, so no test ever touches the participant's database. "
              "Reproducibility follows the recommendations of [@sandve2013]: every figure is generated from an "
              "export file, with the plotted numbers in a CSV beside it, and the figure manifest refits the active "
              f"model from the export and records the largest absolute weight difference "
              f"({d['manifest']['reproducibility_check']['max_abs_weight_difference']:g} for the case-study model)."))
    add(("table", "tests", "Automated test suites at the time of writing.",
         ["Suite", "Tests", "Main properties covered"],
         [["Backend (pytest)", "88", "recomputation and tamper refusal; idempotency; readiness refusals; symmetry; LOSO "
           "isolation; artifact checksums; candidate verification; evaluation protocol; sealed export; figures"],
          ["Unity EditMode", "130", "generator reference values; metrics; walk distances; cost table; validator; "
           "undo/redo; candidate generator; evaluation pair rule; local store and outbox"],
          ["Unity PlayMode", "24 (+5 opt-in)", "offline play; live sync against a throwaway backend; Learning Lab "
           "training; a complete blinded test with all seals verified and a reproducible report"]],
         [0.2, 0.13, 0.67]))

    # 9 Results
    add(("h1", "Case study", "results"))
    add(("h2", "Participant and data", None))
    add(("p", f"The case study involves a single participant, the author, working with synthetic cities. The active "
              f"model ({m['model_version']}) was trained on {n_train} unique comparisons across "
              f"{m['scenario_count']} scenarios (#tab:data#). <b>{n_user} of these were made by the participant</b> "
              f"in the demonstration city (seed 42). <b>The remaining {n_rule} were generated by a scripted agent "
              f"following a fixed rule</b>, \u201cprefer more park access, then lower cost\u201d, in four further "
              f"cities, to reach the five scenarios recommended before a model is called non-exploratory. They are "
              f"flagged in the data, and every figure states that {d['all_rule']} of the {n_all} stored choices were "
              f"rule-made. This composition matters for every result below."))
    rows = []
    for sid in sorted({k[0] for k in ps}, key=lambda s: (s != "green-42", s)):
        nu, nr = ps.get((sid, False), 0), ps.get((sid, True), 0)
        rows.append([sid, str(nu + nr), "participant" if nu else "fixed rule"])
    rows.append(["Total", str(n_train), f"{n_user} participant, {n_rule} rule"])
    add(("table", "data", f"Training data of the active model (A chosen {m['choices_for_a']}, B chosen "
                          f"{m['choices_for_b']}; no comparison ignored as uninformative).",
         ["Scenario", "Comparisons", "Source"], rows, [0.35, 0.25, 0.40]))
    add(("fig", "fig01", FIGS / "fig01_training_data.png", "Training comparisons per scenario (left) and stored "
         "decisions that were never trained on (right): one infeasible pair, the answers from two blinded tests and "
         "three ties.", 1.0))
    add(("h2", "Fitted model", None))
    wrows = []
    for f, label in (("delta_access_fraction", "Park access"), ("delta_housing_fraction", "Housing capacity"),
                     ("delta_park_share_fraction", "Park share of land"), ("negative_budget_fraction", "Low spending"),
                     ("negative_lost_housing_fraction", "Keeping original housing")):
        c = coef[f]
        wrows.append([label, f"{c['weight_standardised']:+.2f}", f"[{c['interval_low']:+.2f}, {c['interval_high']:+.2f}]",
                      f"{c['sign_agreement'] * 100:.0f}%", f"{c['weight_raw']:+.2f}"])
    add(("p", f"#tab:weights# and #fig:coef# show the weights. Park access dominates (standardised weight "
              f"{acc['weight_standardised']:.2f}, 90% interval {acc['interval_low']:.2f} to {acc['interval_high']:.2f}, "
              f"same sign in all resamples). Housing capacity and low spending carry negative weights: chosen designs "
              f"tended to have less housing and higher cost. Park share is near zero and the direction of the "
              f"original-housing weight is uncertain."))
    add(("table", "weights", "Weights of the active model with 90% scenario-bootstrap intervals (200 resamples).",
         ["Feature", "Standardised", "90% interval", "Same sign", "Raw"], wrows, [0.32, 0.17, 0.23, 0.13, 0.15]))
    add(("fig", "coef", FIGS / "fig03_coefficients.png", "Standardised weights with 90% scenario-bootstrap intervals "
         "and sign agreement. Grey intervals include zero.", 0.95))
    add(("p", f"Leave-one-scenario-out accuracy was {pct(diag['model']['accuracy'])} (Brier "
              f"{diag['model']['brier']:.3f}, log loss {diag['model']['log_loss']:.3f}), identical in accuracy to the "
              f"rule \u201calways more park access\u201d ({pct(lrule['more_access']['accuracy'])}, Brier "
              f"{lrule['more_access']['brier']:.3f}) and well above chance ({pct(lrule['chance']['accuracy'])}), the "
              f"equal-weight rule ({pct(lrule['equal_weight']['accuracy'])}) and the cheaper-design rule "
              f"({pct(lrule['cheaper']['accuracy'])}) (#fig:loso#). Training-set accuracy was "
              f"{pct(m['training_fit']['accuracy'])}. The learning curve (#fig:curve#) shows the model approaching, "
              f"but never exceeding, the more-access rule as comparisons accumulate."))
    add(("fig", "loso", FIGS / "fig05_loso_vs_baselines.png", "Leave-one-scenario-out diagnostic of the active model "
         "against four rules (inside the training data; not a held-out test).", 1.0))
    add(("fig", "curve", FIGS / "fig07_learning_curve.png", "Recomputed learning curve: LOSO accuracy of models fitted on "
         "the first n comparisons, with the rules evaluated on the same comparisons.", 0.82))
    add(("h2", "Blinded held-out test", None))
    add(("p", f"The participant ran one blinded test (plan {plan['plan_id']}): {plan['pairs_per_scenario']} pairs in "
              f"each of {len(plan['evaluation_scenarios'])} unseen scenarios plus {plan['repeat_count']} repeated pairs, "
              f"all {uc['answered']} answered. {uc['primary_ties']} first presentations were answered as ties, leaving "
              f"{uc['scored']} scored choices. All {uc['revealed']} seals were opened and verified, the held-out "
              f"guarantee recomputed zero evaluation choices and scenarios in the training data, and the stored report "
              f"re-hashed and recomputed to the same SHA-256 (reproducible = {str(uver['reproducible']).lower()})."))
    rrows = [["Frozen model", f"{um['correct']:g} / {um['n']}", pct(um["accuracy"]),
              f"[{pct(um['accuracy_ci'][0], 0)}, {pct(um['accuracy_ci'][1], 0)}]", f"{um['brier']:.3f}", f"{um['log_loss']:.3f}", "\u2013"]]
    for name in ("more_access", "left_position", "chance", "equal_weight", "cheaper"):
        b = ub[name]
        rrows.append([b["label"].replace(" (always 50 %)", " (0.5)"), f"{b['correct']:g} / {b['n']}", pct(b["accuracy"]),
                      f"[{pct(b['accuracy_ci'][0], 0)}, {pct(b['accuracy_ci'][1], 0)}]", f"{b['brier']:.3f}",
                      f"{b['log_loss']:.3f}" if b["has_log_loss"] else "\u2013",
                      f"{b['model_minus_baseline_accuracy'] * 100:+.1f} [{b['model_minus_baseline_accuracy_ci'][0] * 100:+.0f}, "
                      f"{b['model_minus_baseline_accuracy_ci'][1] * 100:+.0f}]"])
    add(("table", "eval", "Blinded held-out results on the participant's scored choices (95% scenario-bootstrap intervals, "
                          "1000 resamples). The last column is model minus rule accuracy in percentage points.",
         ["Predictor", "Correct", "Accuracy", "95% interval", "Brier", "Log loss", "Model \u2212 rule"],
         rrows, [0.27, 0.09, 0.11, 0.15, 0.08, 0.1, 0.2]))
    add(("p", f"The frozen model predicted {um['correct']:g} of {um['n']} choices ({pct(um['accuracy'])}, 95% interval "
              f"{pct(um['accuracy_ci'][0])} to {pct(um['accuracy_ci'][1])}; #tab:eval#, #fig:evalres#). The rule "
              f"\u201calways more park access\u201d achieved exactly the same, and the two made the same prediction "
              f"on {agree_rule} of the {len(scored)} scored pairs. Per scenario, the model was correct on "
              + ", ".join(f"{s['correct']:g} of {s['scored']}" for s in u["per_scenario"]) +
              f" scored choices. All predictions were confident (favoured-design probability at least "
              f"{pmin:.2f}), so every choice fell in the top calibration bin, where the mean confidence was "
              f"{ucal['bins'][-1]['mean_confidence']:.3f} against an observed rate of "
              f"{ucal['bins'][-1]['observed_rate']:.3f} (ECE {ucal['expected_calibration_error']:.3f}). The single "
              f"error was made with high confidence (#fig:sealed#), which is what drives the log loss "
              f"({um['log_loss']:.3f}) above that of chance ({ub['chance']['log_loss']:.3f}) even though the Brier "
              f"score ({um['brier']:.3f}) is far below chance (0.250)."))
    add(("fig", "evalres", FIGS / f"fig08_evaluation_vs_baselines{tag}.png", "Blinded held-out accuracy, Brier score "
         "and log loss of the frozen model against chance and four rules, with 95% scenario-bootstrap intervals.", 1.0))
    add(("p", f"The participant chose the same design in {urc['agreements']} of {urc['answered_both']} repeated pairs "
              f"({pct(urc['rate'], 0)}, Wilson 95% interval {pct(urc['rate_ci'][0], 0)} to {pct(urc['rate_ci'][1], 0)}), "
              f"and chose the left design in {upos['chose_left']} of {upos['n']} binary choices (two-sided binomial "
              f"p = {upos['binomial_p_value']:.2f}), giving no evidence of a side habit (#fig:consistency#). Because "
              f"the frozen model assigns both presentations of a repeated pair the same probability, its own repeat "
              f"consistency is 100% by construction; the participant's consistency is an approximate ceiling on how "
              f"well any model of these choices can do."))
    add(("fig", "consistency", FIGS / f"fig12_consistency_and_position{tag}.png", "Repeat consistency (left) and the "
         "screen position of the chosen design (right) in the participant's blinded test.", 0.9))
    add(("fig", "sealed", FIGS / f"fig13_sealed_predictions{tag}.png", "Every sealed prediction, revealed after the "
         "choice: the probability the frozen model gave the design the participant chose, in the order shown "
         "(circles: first presentations; squares: repeats; grey: ties).", 0.9))
    if ag:
        am = ag["metrics"]["model"]
        abm = {b["name"]: b for b in ag["metrics"]["baselines"]}
        add(("p", f"A second blinded test was answered by the scripted agent following the same rule that generated "
                  f"the {n_rule} rule-made training comparisons. The model predicted {am['correct']:g} of {am['n']} "
                  f"scored choices, and so did the more-access rule ({pct(abm['more_access']['accuracy'])}). This "
                  f"test mainly confirms that the pipeline works end to end; it is not evidence about a person."))

    # 10 Discussion
    add(("h1", "Discussion", "discussion"))
    add(("p", "<b>The model has not shown more than a one-line rule.</b> On the participant's held-out choices the "
              "frozen model and the rule \u201calways choose more park access\u201d were equally accurate and made the "
              "same prediction on every scored pair. With no discordant pairs, the data carry no information about "
              "which of the two is better, in the sense of McNemar's test for paired proportions [@mcnemar1947]. "
              "This is the expected outcome when one attribute dominates the choices that generated the training "
              "data, and two-thirds of those choices were generated by a rule whose first criterion was exactly "
              "that attribute. A linear model then collapses towards the dominant attribute; to show that it has "
              "learned more, a test needs pairs in which park access conflicts with the other objectives strongly "
              "enough that a person might give it up. The model-blind trade-off rule guarantees conflicting pairs, "
              "but not that the conflicts are large."))
    add(("p", "<b>Weights are associations, not motives.</b> The negative weights on low spending and housing "
              f"capacity do not mean that the participant likes spending or dislikes housing. Creating a park costs "
              f"credits and often replaces buildings, so in the compared pairs more access usually came with higher "
              f"cost and less housing (see the chosen-minus-rejected differences in #fig:tradeoff#). A model that "
              f"tracks access inherits those correlations as weights. The Learning Lab therefore phrases every "
              f"weight as \u201cyour choices were associated with \u2026\u201d and never displays an exchange rate "
              f"built from a weight whose interval includes zero."))
    add(("fig", "tradeoff", FIGS / "fig02_tradeoff_space.png", "Chosen minus rejected design for every training "
         "comparison: park access against housing capacity (left) and against low spending (right). Chosen designs "
         "almost always had more access, often at higher cost.", 1.0))
    add(("p", "<b>Small samples give wide intervals.</b> With nine scored choices, one choice moves accuracy by 11.1 "
              "percentage points, and the 95% interval runs from 50% to 100%, so the test cannot distinguish the model "
              "from guessing at conventional confidence. A percentile bootstrap over only three scenarios is, "
              "moreover, rough and can be optimistic. As a rough planning figure, a normal approximation suggests "
              "the number of scored choices needed for a given interval half-width h:"))
    add(eq("n", r"n \approx \dfrac{z^2\,p(1 - p)}{h^2} = \dfrac{1.96^2 \times 0.85 \times 0.15}{0.10^2} \approx 49"))
    add(("p", "that is, about 50 scored choices for a half-width of 10 points around an accuracy of 85%, and many more "
              "discordant pairs than the zero observed here before the model could be compared with the rule."))
    add(("p", "<b>What the protocol does provide.</b> Within these limits, the result is trustworthy in a specific "
              "sense: the model could not have been trained on the test cities or choices (checked by id, seed and "
              "mode, and recomputed in the report); the predictions demonstrably existed before the choices (seals "
              "verified); the test pairs were not chosen by or for the model; neither the participant nor the "
              "client chose the screen sides; the participant saw no prediction during the test; and the reported "
              "numbers can be recomputed from the stored rows by anyone with the export. These properties are "
              "enforced by refusals in the server, which makes them testable, and they were tested. The same "
              "protocol can evaluate future models (for example, sequence models of the editing process) on equal "
              "terms."))
    add(("p", "<b>Teaching use.</b> For students, the value lies less in the model than in the loop: designing two "
              "complete alternatives under explicit constraints, seeing every metric recomputed transparently, "
              "committing to a choice, and then confronting an honest test of whether a simple summary of one's "
              "choices predicts the next ones, including the possibility that one's choices are less consistent "
              "than expected."))

    # 11 Limitations
    add(("h1", "Limitations and ethical considerations", "limits"))
    add(("bullets", [
        "<b>One participant, who is also the developer.</b> Nothing here generalises to other people, and the "
        "participant's familiarity with the metrics may make choices more rule-like than a student's would be.",
        f"<b>Rule-generated training data.</b> {n_rule} of the {n_train} training comparisons were generated by a "
        "fixed rule. The fitted weights largely reflect that rule, and the agreement between model and rule in the "
        "participant's test is partly inherited from it.",
        "<b>Synthetic cities and proxy metrics.</b> Housing units are a capacity proxy, credits are fictional and "
        "access is a grid-road distance. The system knows nothing about real places, residents, land values or "
        "equity, and its outputs are not planning advice.",
        "<b>Stated, hypothetical choices.</b> Choices between on-screen designs need not reflect what a person "
        "would support in a real decision with real stakes.",
        "<b>Small evaluation.</b> Nine scored choices in three scenarios; intervals are wide and the scenario "
        "bootstrap with three clusters is crude.",
        "<b>Linear model on five features.</b> Interactions, thresholds (for example, caring about access only "
        "below some level) and the reasons participants type are not modelled.",
        "<b>Demonstration video.</b> The accompanying video was recorded on a copy of the data with scripted "
        "choices, and says so; the case-study numbers in this paper come only from the participant's own test.",
    ]))
    add(("p", "All data are synthetic or produced by the author on a single computer; no data about other people were "
              "collected, and nothing leaves the machine. The client offers export and complete deletion of local and "
              "server data. If the system were used with students, participation, data retention and the framing of "
              "model outputs as associations rather than judgments would need appropriate consent and review."))

    # 12 Conclusion
    add(("h1", "Conclusion", "conclusion"))
    add(("p", "PlannerDNA makes the trade-offs of a small planning exercise explicit, records which ones a person "
              "chooses, summarises those choices with a model that can be read line by line, and tests that summary "
              "under a protocol in which the important safeguards (frozen model, unseen scenarios, sealed predictions, "
              "server-held randomisation, model-blind pairs, delayed reveal and a self-verifying report) are enforced "
              f"by the software. In the case study the model predicted {um['correct']:g} of {um['n']} held-out choices, "
              f"exactly as many as the rule \u201calways choose more park access\u201d, and the system reports that "
              f"plainly. We regard that sentence, rather than the accuracy figure, as the main result: a workflow "
              f"that keeps a model honest is worth more than a model that looks good."))

    add(("h1u", "Data and software availability", None))
    add(("p", "All records are stored locally and can be exported as a single JSON file containing scenarios, "
              "sessions, events, alternatives, comparisons, model runs, predictions, evaluation plans and trials. "
              "Figures are regenerated from an export with backend/tools/make_figures.py, which writes the CSV behind "
              "each figure and a manifest with a refit check. Model files are JSON. The scripts that produced the "
              "demonstration video, the carousel and this paper are in the tools folder of the project."))
    add(("h1u", "Use of AI assistance", None))
    add(("p", "Software development and drafting of this report were assisted by an AI coding assistant. All numbers "
              "reported here were produced by the system's own code from the stored data and are pulled from the "
              "local server when the paper is built."))
    add(("refs",))
    return abstract, keywords, B


# ---------------------------------------------------------------- numbering


def number(blocks: list[tuple]):
    sec, sub = 0, 0
    labels: dict[str, str] = {}
    fig_n = tab_n = eq_n = 0
    out = []
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            sec += 1
            sub = 0
            if b[2]:
                labels[f"sec:{b[2]}"] = f"Section {sec}"
            out.append(("h1", f"{sec}\u00a0\u00a0{b[1]}"))
        elif kind == "h2":
            sub += 1
            out.append(("h2", f"{sec}.{sub}\u00a0\u00a0{b[1]}"))
        elif kind == "fig":
            fig_n += 1
            labels[f"fig:{b[1]}"] = f"Fig. {fig_n}"
            out.append(("fig", fig_n) + b[1:])
        elif kind == "table":
            tab_n += 1
            labels[f"tab:{b[1]}"] = f"Table {tab_n}"
            out.append(("table", tab_n) + b[1:])
        elif kind == "eq":
            eq_n += 1
            labels[f"eq:{b[1]}"] = f"({eq_n})"
            out.append(("eq", eq_n) + b[1:])
        else:
            out.append(b)
    return out, labels


def cite_and_ref(blocks, labels, abstract):
    order: list[str] = []

    def nums(keys):
        ns = []
        for k in keys:
            if k not in REFS:
                raise SystemExit(f"Unknown reference {k}")
            if k not in order:
                order.append(k)
            ns.append(order.index(k) + 1)
        ns = sorted(set(ns))
        parts, i = [], 0
        while i < len(ns):
            j = i
            while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                j += 1
            parts.append(f"{ns[i]}\u2013{ns[j]}" if j - i >= 2 else ", ".join(str(n) for n in ns[i:j + 1]))
            i = j + 1
        return "[" + ", ".join(parts) + "]"

    def sub(text: str) -> str:
        text = re.sub(r"\[((?:@[\w]+(?:;\s*)?)+)\]", lambda mm: nums(re.findall(r"@(\w+)", mm.group(1))), text)
        text = re.sub(r"#((?:sec|fig|tab|eq):\w+)#", lambda mm: labels[mm.group(1)], text)
        return text

    out = []
    for b in blocks:
        if b[0] == "p":
            out.append(("p", sub(b[1])))
        elif b[0] == "bullets":
            out.append(("bullets", [sub(t) for t in b[1]]))
        elif b[0] == "fig":
            out.append(b[:4] + (sub(b[4]),) + b[5:])
        elif b[0] == "table":
            out.append(b[:3] + (sub(b[3]), b[4], [[sub(c) for c in r] for r in b[5]], b[6]))
        else:
            out.append(b)
    return out, order, sub(abstract)


# ---------------------------------------------------------------- PDF


def build_pdf(path: Path, abstract: str, keywords: list[str], blocks, order) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm, inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, KeepTogether, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.platypus import Paragraph as _Paragraph
    from reportlab.lib.fonts import addMapping

    for name, file in (("TNR", "times.ttf"), ("TNR-B", "timesbd.ttf"), ("TNR-I", "timesi.ttf"), ("TNR-BI", "timesbi.ttf")):
        pdfmetrics.registerFont(TTFont(name, str(WIN_FONTS / file)))
    addMapping("TNR", 0, 0, "TNR")
    addMapping("TNR", 1, 0, "TNR-B")
    addMapping("TNR", 0, 1, "TNR-I")
    addMapping("TNR", 1, 1, "TNR-BI")

    body = ParagraphStyle("body", fontName="TNR", fontSize=10.5, leading=13.6, alignment=TA_JUSTIFY, spaceAfter=6)
    h1 = ParagraphStyle("h1", fontName="TNR-B", fontSize=12.5, leading=16, spaceBefore=12, spaceAfter=6, keepWithNext=1)
    h2 = ParagraphStyle("h2", fontName="TNR-BI", fontSize=10.8, leading=14, spaceBefore=8, spaceAfter=4, keepWithNext=1)
    cap = ParagraphStyle("cap", fontName="TNR", fontSize=9, leading=11.2, alignment=TA_JUSTIFY)
    cell = ParagraphStyle("cell", fontName="TNR", fontSize=8.6, leading=10.4)
    cellb = ParagraphStyle("cellb", parent=cell, fontName="TNR-B")
    ref = ParagraphStyle("ref", fontName="TNR", fontSize=9, leading=11.2, leftIndent=0.85 * cm, bulletIndent=0, spaceAfter=3)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=0.6 * cm, bulletIndent=0.15 * cm, spaceAfter=3)

    W, Hh = A4
    margin = 2.2 * cm
    tw = W - 2 * margin

    def Paragraph(text, style, **kw):  # noqa: N802 - shadows reportlab's class with sub/sup metrics
        s = style.fontSize
        text = text.replace("<sub>", f'<sub rise="{0.2 * s:.1f}" size="{0.72 * s:.1f}">')
        text = text.replace("<sup>", f'<sup rise="{0.36 * s:.1f}" size="{0.72 * s:.1f}">')
        return _Paragraph(text, style, **kw)

    def on_page(c, doc):
        c.saveState()
        c.setFont("TNR", 8.5)
        c.setFillColor(colors.HexColor("#475569"))
        c.drawCentredString(W / 2, 1.2 * cm, str(doc.page))
        if doc.page > 1:
            c.drawString(margin, Hh - 1.3 * cm, SHORT_TITLE)
            c.setStrokeColor(colors.HexColor("#cbd5e1"))
            c.setLineWidth(0.5)
            c.line(margin, Hh - 1.45 * cm, W - margin, Hh - 1.45 * cm)
        c.restoreState()

    story = [
        Paragraph(TITLE, ParagraphStyle("t", fontName="TNR-B", fontSize=17, leading=21, alignment=TA_CENTER, spaceAfter=10)),
        Paragraph(AUTHOR, ParagraphStyle("a", fontName="TNR", fontSize=11.5, leading=14, alignment=TA_CENTER)),
        Paragraph(f"<i>{AFFILIATION}</i> \u2022 Technical report \u2022 {DATE}",
                  ParagraphStyle("af", fontName="TNR", fontSize=10, leading=13, alignment=TA_CENTER, spaceAfter=14)),
    ]
    abs_style = ParagraphStyle("abs", parent=body, fontSize=9.8, leading=12.6, leftIndent=0.8 * cm, rightIndent=0.8 * cm)
    story.append(Paragraph(f"<b>Abstract.</b> {abstract}", abs_style))
    story.append(Paragraph(f"<b>Keywords:</b> {'; '.join(keywords)}.", ParagraphStyle("kw", parent=abs_style, spaceAfter=10)))

    for b in blocks:
        kind = b[0]
        if kind == "h1":
            story.append(Paragraph(b[1], h1))
        elif kind == "h1u":
            story.append(Paragraph(b[1], h1))
        elif kind == "h2":
            story.append(Paragraph(b[1], h2))
        elif kind == "p":
            story.append(Paragraph(b[1], body))
        elif kind == "bullets":
            for t in b[1]:
                story.append(Paragraph(t, bullet, bulletText="\u2022"))
            story.append(Spacer(1, 3))
        elif kind == "eq":
            n, _key, p, (wi, hi) = b[1], b[2], b[3], b[4]
            scale = min(1.0, (tw - 1.6 * cm) / (wi * inch))
            img = Image(str(p), width=wi * inch * scale, height=hi * inch * scale)
            t = Table([[img, Paragraph(f"({n})", ParagraphStyle("en", parent=body, alignment=2))]],
                      colWidths=[tw - 1.4 * cm, 1.4 * cm])
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "CENTER"),
                                   ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
            story.append(t)
        elif kind == "fig":
            n, _key, p, caption, frac = b[1], b[2], b[3], b[4], b[5]
            from PIL import Image as PImage
            with PImage.open(p) as im:
                iw, ih = im.size
            w = tw * frac
            h = w * ih / iw
            if h > 15.5 * cm:
                h = 15.5 * cm
                w = h * iw / ih
            img = Image(str(p), width=w, height=h)
            story.append(KeepTogether([Spacer(1, 4), img, Spacer(1, 4),
                                       Paragraph(f"<b>Fig. {n}.</b> {caption}", cap), Spacer(1, 8)]))
        elif kind == "table":
            n, _key, caption, header, rows, widths = b[1:]
            data = [[Paragraph(h, cellb) for h in header]] + [[Paragraph(c, cell) for c in r] for r in rows]
            t = Table(data, colWidths=[tw * f for f in widths], repeatRows=1)
            t.setStyle(TableStyle([
                ("LINEABOVE", (0, 0), (-1, 0), 1.0, colors.black),
                ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
                ("LINEBELOW", (0, -1), (-1, -1), 1.0, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ]))
            story.append(KeepTogether([Paragraph(f"<b>Table {n}.</b> {caption}", ParagraphStyle("tc", parent=cap, spaceAfter=4)),
                                       t, Spacer(1, 10)]))
        elif kind == "refs":
            story.append(Paragraph("References", h1))
            for i, k in enumerate(order, 1):
                story.append(Paragraph(REFS[k], ref, bulletText=f"[{i}]"))

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=2.0 * cm,
                            bottomMargin=2.0 * cm, title=TITLE, author=AUTHOR, subject="Technical report")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


# ---------------------------------------------------------------- DOCX

_TAG = re.compile(r"(<b>|</b>|<i>|</i>|<sub>|</sub>|<sup>|</sup>)")


def _runs(par, text: str, size=None, base_bold=False, base_italic=False):
    from docx.shared import Pt
    state = {"b": base_bold, "i": base_italic, "sub": False, "sup": False}
    for tok in _TAG.split(text):
        if not tok:
            continue
        m = re.fullmatch(r"<(/?)(b|i|sub|sup)>", tok)
        if m:
            state[m.group(2)] = not m.group(1)
            if m.group(1) and m.group(2) in ("b", "i"):
                state[m.group(2)] = base_bold if m.group(2) == "b" else base_italic
            continue
        r = par.add_run(html.unescape(tok))
        r.bold, r.italic = state["b"], state["i"]
        r.font.subscript, r.font.superscript = state["sub"], state["sup"]
        if size:
            r.font.size = Pt(size)


def _set_font(style, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(a), None)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = RGBColor.from_string(color or "000000")


def _no_borders(table):
    from docx.oxml.ns import qn
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = borders.makeelement(qn(f"w:{edge}"), {qn("w:val"): "nil"})
        borders.append(e)
    tbl_pr.append(borders)


def _booktabs(table, n_rows):
    from docx.oxml.ns import qn
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge, val, sz in (("top", "single", "12"), ("bottom", "single", "12"), ("left", "nil", "0"),
                          ("right", "nil", "0"), ("insideH", "nil", "0"), ("insideV", "nil", "0")):
        borders.append(borders.makeelement(qn(f"w:{edge}"), {qn("w:val"): val, qn("w:sz"): sz, qn("w:color"): "000000"}))
    tbl_pr.append(borders)
    for c in table.rows[0].cells:
        tc_pr = c._tc.get_or_add_tcPr()
        b = tc_pr.makeelement(qn("w:tcBorders"), {})
        b.append(b.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "6", qn("w:color"): "000000"}))
        tc_pr.append(b)


def _page_field(par):
    from docx.oxml.ns import qn
    run = par.add_run()
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "end"})):
        el = run._r.makeelement(qn(tag), {qn(k): v for k, v in (attr or {}).items()})
        if tag == "w:instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        run._r.append(el)


def build_docx(path: Path, abstract: str, keywords: list[str], blocks, order) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Inches, Pt
    from PIL import Image as PImage

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    tw_cm = 21.0 - 4.4
    st = doc.styles
    _set_font(st["Normal"], size=11)
    st["Normal"].paragraph_format.space_after = Pt(6)
    st["Normal"].paragraph_format.line_spacing = 1.15
    _set_font(st["Heading 1"], size=13, bold=True)
    _set_font(st["Heading 2"], size=11.5, bold=True, italic=True)
    for h in ("Heading 1", "Heading 2"):
        st[h].paragraph_format.space_before = Pt(12 if h == "Heading 1" else 8)
        st[h].paragraph_format.space_after = Pt(4)
        st[h].paragraph_format.keep_with_next = True
    _set_font(st["Title"], size=17, bold=True)

    hp = sec.header.paragraphs[0]
    _runs(hp, SHORT_TITLE, size=8.5, base_italic=True)
    sec.different_first_page_header_footer = True
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_field(fp)
    fp2 = sec.first_page_footer.paragraphs[0]
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_field(fp2)

    t = doc.add_paragraph(style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(t, TITLE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(p, AUTHOR, size=12)
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(p, f"<i>{AFFILIATION}</i> \u2022 Technical report \u2022 {DATE}", size=10.5)
    p.paragraph_format.space_after = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(0.8)
    _runs(p, f"<b>Abstract.</b> {abstract}", size=10)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(12)
    _runs(p, f"<b>Keywords:</b> {'; '.join(keywords)}.", size=10)

    def caption(text, before=False):
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cp.paragraph_format.space_after = Pt(4 if before else 10)
        if before:
            cp.paragraph_format.keep_with_next = True
        _runs(cp, text, size=9)

    for b in blocks:
        kind = b[0]
        if kind in ("h1", "h1u"):
            hd = doc.add_paragraph(style="Heading 1")
            _runs(hd, b[1])
        elif kind == "h2":
            hd = doc.add_paragraph(style="Heading 2")
            _runs(hd, b[1])
        elif kind == "p":
            pp = doc.add_paragraph()
            pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _runs(pp, b[1])
        elif kind == "bullets":
            for item in b[1]:
                pp = doc.add_paragraph(style="List Bullet")
                pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                pp.paragraph_format.space_after = Pt(3)
                _runs(pp, item)
        elif kind == "eq":
            n, _key, pth, (wi, hi) = b[1], b[2], b[3], b[4]
            tb = doc.add_table(rows=1, cols=2)
            _no_borders(tb)
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            c0, c1 = tb.rows[0].cells
            c0.width, c1.width = Cm(tw_cm - 1.5), Cm(1.5)
            ip = c0.paragraphs[0]
            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.add_run().add_picture(str(pth), width=Inches(min(wi, (tw_cm - 1.8) / 2.54)))
            np_ = c1.paragraphs[0]
            np_.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _runs(np_, f"({n})")
            from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
            c0.vertical_alignment = c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        elif kind == "fig":
            n, _key, pth, cap_text, frac = b[1], b[2], b[3], b[4], b[5]
            with PImage.open(pth) as im:
                iw, ih = im.size
            wcm = tw_cm * frac
            if wcm * ih / iw > 15.5:
                wcm = 15.5 * iw / ih
            ip = doc.add_paragraph()
            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.paragraph_format.keep_with_next = True
            ip.add_run().add_picture(str(pth), width=Cm(wcm))
            caption(f"<b>Fig. {n}.</b> {cap_text}")
        elif kind == "table":
            n, _key, cap_text, header, rows, widths = b[1:]
            caption(f"<b>Table {n}.</b> {cap_text}", before=True)
            tb = doc.add_table(rows=1 + len(rows), cols=len(header))
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            tb.autofit = False
            _booktabs(tb, len(rows) + 1)
            from docx.oxml.ns import qn
            layout = tb._tbl.tblPr.makeelement(qn("w:tblLayout"), {qn("w:type"): "fixed"})
            tb._tbl.tblPr.append(layout)
            for ci, col in enumerate(tb.columns):
                col.width = Cm(tw_cm * widths[ci])
            for ri, r in enumerate([header] + rows):
                for ci, text in enumerate(r):
                    c = tb.rows[ri].cells[ci]
                    c.width = Cm(tw_cm * widths[ci])
                    cp = c.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(1)
                    cp.paragraph_format.line_spacing = 1.0
                    _runs(cp, text, size=9, base_bold=ri == 0)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "refs":
            hd = doc.add_paragraph(style="Heading 1")
            _runs(hd, "References")
            for i, k in enumerate(order, 1):
                rp = doc.add_paragraph()
                rp.paragraph_format.left_indent = Cm(0.9)
                rp.paragraph_format.first_line_indent = Cm(-0.9)
                rp.paragraph_format.space_after = Pt(3)
                _runs(rp, f"[{i}]\t{REFS[k]}", size=9.5)
    core = doc.core_properties
    core.title, core.author, core.subject = TITLE, AUTHOR, "Technical report"
    doc.save(str(path))


# ---------------------------------------------------------------- main


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--api", default="http://127.0.0.1:8765")
    ap.add_argument("--out", default=str(DELIV / "PlannerDNA_Technical_Paper"))
    args = ap.parse_args(argv)

    plt.rcParams["font.family"] = "Arial"
    ASSETS.mkdir(parents=True, exist_ok=True)
    eqdir = ASSETS / "equations"
    eqdir.mkdir(exist_ok=True)
    diagrams = {"arch": diagram_architecture, "label": diagram_labelling, "train": diagram_training,
                "protocol": diagram_protocol, "walk": diagram_walk}
    paths = {}
    for key, fn in diagrams.items():
        paths[key] = ASSETS / f"diagram_{key}.png"
        fn(paths[key])

    data = load(args.api)
    abstract, keywords, blocks = build_content(data, eqdir)
    blocks = [(b[0], b[1], paths[b[1]]) + b[3:] if b[0] == "fig" and b[2] is None else b for b in blocks]
    blocks, labels = number(blocks)
    blocks, order, abstract = cite_and_ref(blocks, labels, abstract)

    out = Path(args.out)
    build_pdf(out.with_suffix(".pdf"), abstract, keywords, blocks, order)
    build_docx(out.with_suffix(".docx"), abstract, keywords, blocks, order)
    print(f"Wrote {out.with_suffix('.pdf')} and {out.with_suffix('.docx')}: "
          f"{sum(1 for b in blocks if b[0] == 'fig')} figures, {sum(1 for b in blocks if b[0] == 'table')} tables, "
          f"{sum(1 for b in blocks if b[0] == 'eq')} equations, {len(order)} references")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
